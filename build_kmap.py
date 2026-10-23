# -*- coding: utf-8 -*-
"""「能力图谱」页面：读取 data_kmap.py 的专属数据，产出树形图所需 JSON，注入 kmap_template.html。

data_kmap.SKILL_SOURCE == 'knowledge_system' 时，技能层由专业知识体系 + 能力↔技能映射生成（文件与口径见 data_assess.py）：
能力项 → 技能规范（已确认的直接 / 部分支撑映射）→ 操作步骤 / 标准 / 案例 / 知识点；节点 ID 按知识体系路径生成，与能力测评一致。
"""
import os, sys, json, re, hashlib
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import data_kmap as K

LV = {"L1": 1, "L2": 2, "L3": 3}
TYPE_CN = {"knowledge": "知识", "skill": "技能", "literacy": "素养"}

subs = [{"id": i, "name": n} for i, n in K.SUB_INDUSTRIES]
positions = [{"id": i, "name": n, "sub": s, "aliases": al} for i, n, s, al in K.POSITIONS]
abilities = [{"id": i, "name": n, "type": t, "domain": d, "core": c} for i, n, t, d, c in K.ABILITIES]
CHILD_KINDS = ("知识点", "操作步骤", "标准", "案例")

def nid(prefix, path): return prefix + hashlib.sha1(path.encode("utf-8")).hexdigest()[:10]

def from_knowledge_system():
    """专业知识体系 + 能力↔技能映射 → SKILLS / ABILITY_SKILL / SKILL_NODES / SKILL_CHILD，并给每个技能带上知识体系位置。"""
    import openpyxl
    import data_assess as A
    abl_ids = {a[0] for a in K.ABILITIES}
    abl_by_name = {a[1]: a[0] for a in K.ABILITIES}
    SK, AS, SN, SC, KS = [], [], [], [], {}
    seen_nodes = set()
    for mid, f in A.KNOWLEDGE_FILES.items():
        wb = openpyxl.load_workbook(os.path.join(HERE, f), read_only=True)
        level = {}
        for r in list(wb["节点属性数据"].iter_rows(values_only=True))[3:]:
            path, info = (list(r) + [None] * 3)[0], (list(r) + [None] * 3)[2]
            if not path: continue
            kv = {}
            for part in re.split("[；;]", str(info or "")):
                if re.search("[:：]", part):
                    k, v = re.split("[:：]", part, 1); kv[k.strip()] = v.strip()
            level[str(path).strip()] = kv.get("层级") or kv.get("类型")
        kps = defaultdict(list)   # 技能规范路径 → 知识点名称
        for r in list(wb["知识点数据"].iter_rows(values_only=True))[3:]:
            path, name = (list(r) + [None] * 2)[:2]
            if path and name: kps[str(path).strip()].append(str(name).strip())
        skill_paths = [p for p, l in level.items() if l == "技能规范"]
        # 映射：只计已确认、直接 / 部分支撑（演示环境按 data_assess 的设置把草稿视为已确认）
        maps = json.load(open(os.path.join(HERE, A.MAPPING_FILES[mid]), encoding="utf-8"))
        demo = mid in getattr(A, "DEMO_TREAT_DRAFT_AS_CONFIRMED", [])
        for m in maps:
            st = m.get("status", "待审核")
            if demo and st == "待审核": st = "已审核"
            if st not in A.CONFIRMED_STATUSES or m.get("mapping_type") not in A.COUNTED_TYPES: continue
            sp = m.get("knowledge_node_path")
            assert sp in level and level[sp] == "技能规范", f"映射的技能在知识体系中找不到：{m.get('knowledge_node_name')}"
            aid = m.get("ability_node_id") if m.get("ability_node_id") in abl_ids else abl_by_name.get(m.get("ability_node_name"))
            assert aid, f"映射的能力在图谱中找不到：{m.get('ability_node_name')}"
            pair = (aid, nid("K", sp))
            if pair not in AS: AS.append(pair)
        used = {k for _, k in AS}
        for sp in skill_paths:
            kid = nid("K", sp)
            if kid not in used or kid in KS: continue
            seg = sp.split("--")
            SK.append((kid, seg[-1]))
            # 知识体系位置：专业 › 课程分类 › 课程 › 任务（含各级 ID，供带入备课组装 knowledge_system）
            anc = ["--".join(seg[:i]) for i in range(1, len(seg))]
            KS[kid] = [{"id": nid({"专业": "M", "课程分类": "G", "课程": "C", "任务": "T"}.get(level.get(a), "X"), a),
                        "type": level.get(a) or "", "name": a.split("--")[-1]} for a in anc]
            for cp in [p for p in level if p.startswith(sp + "--") and p.count("--") == sp.count("--") + 1]:
                kind = level[cp]
                if kind not in CHILD_KINDS: continue
                cid = nid("N", cp)
                if cid not in seen_nodes: SN.append((cid, kind, cp.split("--")[-1])); seen_nodes.add(cid)
                SC.append((kid, cid))
            for i, name in enumerate(kps.get(sp, [])):
                cid = nid("P", sp + "|" + name + "|" + str(i))
                if cid not in seen_nodes: SN.append((cid, "知识点", name)); seen_nodes.add(cid)
                SC.append((kid, cid))
    return SK, AS, SN, SC, KS

if getattr(K, "SKILL_SOURCE", "") == "knowledge_system":
    SKILLS_SRC, ABILITY_SKILL_SRC, SKILL_NODES_SRC, SKILL_CHILD_SRC, SKILL_KS = from_knowledge_system()
else:
    SKILLS_SRC, ABILITY_SKILL_SRC, SKILL_NODES_SRC, SKILL_CHILD_SRC, SKILL_KS = K.SKILLS, K.ABILITY_SKILL, K.SKILL_NODES, K.SKILL_CHILD, {}
skills = [{"id": i, "name": n} for i, n in SKILLS_SRC]
nodes = [{"id": i, "kind": k, "name": n} for i, k, n in SKILL_NODES_SRC]
for k in skills:
    if k["id"] in SKILL_KS: k["ks"] = SKILL_KS[k["id"]]

# 新兴标记：data_kmap.py 里的 EMERGING_* id 集合，标到对应节点上，供页面显示「新兴」徽标
for s in subs: s["emerging"] = s["id"] in getattr(K, "EMERGING_SUBS", set())
for p in positions: p["emerging"] = p["id"] in getattr(K, "EMERGING_POSITIONS", set())
for a in abilities: a["emerging"] = a["id"] in getattr(K, "EMERGING_ABILITIES", set())
for k in skills: k["emerging"] = k["id"] in getattr(K, "EMERGING_SKILLS", set())
# 新兴标记必须指向存在的节点，且每个都有原因
for name, ids, pool in (("子行业", getattr(K, "EMERGING_SUBS", set()), subs), ("岗位", getattr(K, "EMERGING_POSITIONS", set()), positions),
                        ("能力项", getattr(K, "EMERGING_ABILITIES", set()), abilities), ("技能", getattr(K, "EMERGING_SKILLS", set()), skills)):
    miss = set(ids) - {x["id"] for x in pool}
    assert not miss, f"新兴标记的{name}不存在：{sorted(miss)}"
    nor = set(ids) - set(getattr(K, "EMERGING_REASONS", {}))
    assert not nor, f"新兴{name}缺少原因：{sorted(nor)}"

# 节点详情（NODE_DETAILS）合入各节点 detail 字段，供页面「详情」弹窗展示
ND = getattr(K, "NODE_DETAILS", {})
for s in subs: s["detail"] = ND.get(s["id"], {})
for p in positions: p["detail"] = ND.get(p["id"], {})
for a in abilities: a["detail"] = ND.get(a["id"], {})
industry_detail = ND.get(K.INDUSTRY["id"], {})

abl_by = {a["id"]: a for a in abilities}
pos_by = {p["id"]: p for p in positions}
sub_by = {s["id"]: s for s in subs}
skl_by = {s["id"]: s for s in skills}
node_by = {n["id"]: n for n in nodes}

by_pos, by_abl = defaultdict(list), defaultdict(list)
for p, a, w, lv in K.POSITION_ABILITY:
    assert p in pos_by, f"岗位不存在：{p}"
    assert a in abl_by, f"能力项不存在：{a}"
    assert lv in LV, f"要求等级只能是 L1/L2/L3：{p} {a} {lv}"
    by_pos[p].append({"ability": a, "weight": w, "level": lv})
    by_abl[a].append({"position": p, "weight": w, "level": lv})
for k in by_pos: by_pos[k].sort(key=lambda x: -x["weight"])
for k in by_abl: by_abl[k].sort(key=lambda x: -x["weight"])

for p in positions:
    assert p["sub"] in sub_by, f"岗位 {p['id']} 的子行业不存在：{p['sub']}"
orphan = [a["id"] for a in abilities if not by_abl[a["id"]]]
assert not orphan, f"未被任何岗位引用的能力项：{orphan}"
nopos = [p["id"] for p in positions if not by_pos[p["id"]]]
assert not nopos, f"没有任何能力要求的岗位：{nopos}"

for p in positions:
    rows = by_pos[p["id"]]
    p["abilities"] = rows
    p["abilityCount"] = len(rows)
    p["l3Count"] = sum(1 for r in rows if r["level"] == "L3")

for s in subs:
    ps = [p for p in positions if p["sub"] == s["id"]]
    s["positionIds"] = [p["id"] for p in ps]
    s["distinctAbilityCount"] = len({r["ability"] for p in ps for r in by_pos[p["id"]]})
    s["built"] = len(ps) > 0

for a in abilities:
    rows = by_abl[a["id"]]
    a["requiredBy"] = rows
    a["positionCount"] = len(rows)
    a["maxRequiredLevel"] = "L%d" % max(LV[r["level"]] for r in rows)
    a["typeCn"] = TYPE_CN[a["type"]]
    a["acrossSubIndustries"] = sorted({pos_by[r["position"]]["sub"] for r in rows})

abl_skills, skl_abls = defaultdict(list), defaultdict(list)
for a, k in ABILITY_SKILL_SRC:
    assert a in abl_by, f"能力项不存在：{a}"
    assert k in skl_by, f"技能不存在：{k}"
    assert k not in abl_skills[a], f"重复的能力项 → 技能：{a} {k}"
    abl_skills[a].append(k)
    skl_abls[k].append(a)
for a in abilities:
    a["skills"] = abl_skills[a["id"]]
for k in skills:
    k["abilityIds"] = skl_abls[k["id"]]
unused = [k["id"] for k in skills if not k["abilityIds"]]
assert not unused, f"未归属任何能力项的技能：{unused}"

for n in nodes:
    assert n["kind"] in CHILD_KINDS, f"下级节点类型只能是 {'/'.join(CHILD_KINDS)}：{n['id']} {n['kind']}"
skl_children = defaultdict(list)
for k, n in SKILL_CHILD_SRC:
    assert k in skl_by, f"技能不存在：{k}"
    assert n in node_by, f"下级节点不存在：{n}"
    assert n not in skl_children[k], f"重复的技能 → 下级节点：{k} {n}"
    skl_children[k].append(n)
for k in skills:   # 按「知识点 / 操作步骤 / 标准 / 案例」分组，组内保持数据顺序
    k["children"] = sorted(skl_children[k["id"]], key=lambda n: CHILD_KINDS.index(node_by[n]["kind"]))
used = {n for k in skills for n in k["children"]}
assert used == set(node_by), f"未归属任何技能的下级节点：{sorted(set(node_by) - used)}"

# 自检规则：只在 skill 里维护，构建时取 RULES 标记之间的内容带进页面
skill_path = os.path.join(HERE, ".claude", "skills", "kmap-ai-selfcheck", "SKILL.md")
skill_md = open(skill_path, encoding="utf-8").read()
m = re.search(r"<!-- RULES:START -->(.*?)<!-- RULES:END -->", skill_md, re.S)
assert m, f"{skill_path} 缺少 RULES:START / RULES:END 标记"
selfcheck_rules = m.group(1).strip()

DATA = {"industry": {**dict(K.INDUSTRY), "detail": industry_detail}, "subIndustries": subs,
        "positions": positions, "abilities": abilities, "skills": skills, "skillNodes": nodes,
        "emergingReasons": dict(getattr(K, "EMERGING_REASONS", {})),
        "selfCheckRules": selfcheck_rules}
payload = json.dumps(DATA, ensure_ascii=False, separators=(",", ":"))

# AI 预置配置：ai_config.json 存在则注入（部署时把 Key 写在这里），不存在则空配置（保持现状，需使用者自填）
AI_CFG_PATH = os.path.join(HERE, "ai_config.json")
ai_cfg = {}
if os.path.exists(AI_CFG_PATH):
    try:
        raw_cfg = json.load(open(AI_CFG_PATH, encoding="utf-8"))
        for k in ("type", "url", "model", "key"):
            if k in raw_cfg and isinstance(raw_cfg[k], str):
                ai_cfg[k] = raw_cfg[k]
    except Exception as e:
        print(f"警告：ai_config.json 解析失败，已按空配置处理（{e}）")
ai_cfg_json = json.dumps(ai_cfg, ensure_ascii=False, separators=(",", ":"))

tpl = open(os.path.join(HERE, "kmap_template.html"), encoding="utf-8").read()
assert "/*__DATA__*/" in tpl, "模板缺少数据占位符"
assert "/*__AI_CFG__*/" in tpl, "模板缺少 AI 配置占位符"
out_path = os.path.join(HERE, "kmap.html")
html = tpl.replace("/*__DATA__*/", payload).replace("/*__AI_CFG__*/", ai_cfg_json)
open(out_path, "w", encoding="utf-8").write(html)

print(f"子行业 {sum(1 for s in subs if s['built'])}（共 {len(subs)}）· 岗位 {len(positions)} · "
      f"能力项 {len(abilities)} · 能力要求 {len(K.POSITION_ABILITY)} · "
      f"技能 {len(skills)}（能力项 → 技能 {len(ABILITY_SKILL_SRC)}，无技能的能力项 {sum(1 for a in abilities if not a['skills'])}）· "
      f"下级节点 {len(nodes)}（技能 → 下级节点 {len(SKILL_CHILD_SRC)}）")
print(f"自检规则 {len(selfcheck_rules)} 字（来自 .claude/skills/kmap-ai-selfcheck/SKILL.md）")
print(f"输出 {out_path}（{len(payload)/1024:.0f} KB 数据）")
if ai_cfg.get("key"):
    print(f"AI 预置配置：已注入 {ai_cfg.get('type', '')} / {ai_cfg.get('model', '')}")
else:
    print("AI 预置配置：Key 为空（未提供 ai_config.json 或其中 key 为空），AI 功能需使用者自填 Key")
